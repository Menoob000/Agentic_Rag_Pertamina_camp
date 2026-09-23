import sys
import os
from pathlib import Path
from docx import Document

from tender.docx import render_tender_to_docx

def test_tender_rendering():
    sample_data = {
        "cover_page": {
            "nama_pengadaan": "PENGADAAN DAN PEMASANGAN PIPA TRANSMISI MINYAK DAN GAS REGIONAL KALIMANTAN",
            "nomor_tender": "No.Project/DT/PND970000/2026-S7",
            "tanggal": "04 Desember 2025",
            "pejabat_procurement": {
                "jabatan_1": "Wkl. Fungsi Procurement (Pengadaan)",
                "jabatan_2": "Area Manager Procurement Kalimantan",
                "nama": "Rigga Widar Atmagi"
            },
            "unit_regional": {
                "nama_entitas": "PT PERTAMINA PATRA NIAGA REGIONAL KALIMANTAN",
                "fungsi": "PELAKSANA PEMILIHAN PENYEDIA FUNGSI PROCUREMENT KALIMANTAN",
                "alamat": "Jln. Yos Sudarso No. 148 Balikpapan 76123",
                "telepon": "(0542) 752 4444",
                "website": "www.pertaminapatraniaga.com"
            }
        },
        "ketentuan_khusus_ikpp": {
            "penyelenggara": "Fungsi Procurement (Pengadaan)",
            "kategori_pengadaan": "Pengadaan dalam Kondisi Normal",
            "pemilihan_penyedia": "Pertama",
            "syarat_status_peserta": "Tunggal",
            "syarat_golongan_usaha": "Menengah",
            "syarat_kualifikasi_csms": {
                "tingkat_risiko": "Tinggi",
                "kualifikasi_minimal": "Tinggi"
            },
            "syarat_kbup_kbli": [
                {
                    "kode": "Q - Jasa Pelaksana Konstruksi",
                    "kode_sub_bidang": "Q.13.09",
                    "deskripsi": "Konstruksi Perpipaan Minyak, Gas, dan Energi termasuk perawatannya (Pekerjaan Rekayasa)"
                },
                {
                    "kode": "Q - Jasa Pelaksana Konstruksi",
                    "kode_sub_bidang": "Q.13.10",
                    "deskripsi": "Fasilitas Produksi, Perpipaan Minyak dan Gas, Termasuk Perawatannya (Pekerjaan Rekayasa)"
                }
            ],
            "metode_pemenuhan": "Tender Terbatas",
            "pejabat_berwenang": "Sr. Manager Opt. & Maint. Regional Kalimantan",
            "pengawas_pekerjaan": "Region Manager RPD Regional Kalimantan",
            "jadwal_prebid": {
                "mekanisme": "Online",
                "hari_tanggal": "Senin, 08 Desember 2025",
                "waktu": "10.00 WITA",
                "tempat": "Microsoft Teams Meeting dengan link yang disampaikan melalui email Undangan Prebid Meeting"
            },
            "jadwal_pemasukan": {
                "mulai_hari_tanggal": "Senin, 08 Desember 2025",
                "mulai_waktu": "09.00 WITA",
                "selesai_hari_tanggal": "Senin, 15 Desember 2025",
                "selesai_waktu": "16.00 WITA",
                "tempat_portal": "https://smart.gep.com"
            },
            "masa_berlaku_penawaran": "Penawaran berlaku selama 90 (sembilan puluh) Hari Kalender sejak Pembukaan Dokumen Penawaran",
            "metode_penyampaian_penawaran": "1 (satu) tahap 1 (satu) Sampul",
            "parameter_penetapan_pemenang": "HEA Terbaik",
            "metode_peringkat_peserta": "Peringkat Peserta ditentukan dengan memperhatikan metode HEA Terbaik",
            "evaluasi_teknis": {
                "metode": "Scoring",
                "passing_grade": "80 dari skala 100 (seratus)"
            },
            "evaluasi_hsse_plan": {
                "dipersyaratkan": True,
                "tingkat_risiko": "High Risk",
                "metode": "Scoring",
                "passing_grade": "Minimal 80% dalam skala 100%"
            },
            "evaluasi_tkdn": {
                "minimal_persen": 20.12,
                "komponen_barang": []
            },
            "jenis_kontrak": "Gabungan Harga Satuan & Lumpsum",
            "jumlah_pemenang": "Single Winner",
            "hps_oe": {
                "prime_cost": 13283434534.0,
                "prime_cost_str": "Rp13.283.434.534,00",
                "total_dengan_kr": 14346100000.0,
                "total_dengan_kr_str": "Rp14.346.100.000,00"
            }
        },
        "lampiran_2a_evaluasi_teknis": [
            {
                "no": 1,
                "kriteria": "Metode Kerja & Rencana Pelaksanaan Konstruksi Pipa",
                "dokumen_diperlukan": "Dokumen Metode Kerja, CPM, dan Schedule",
                "parameter_evaluasi": "Kesesuaian metode stringing, trenching, welding, and hydrostatic testing",
                "bobot": 30
            },
            {
                "no": 2,
                "kriteria": "Tenaga Ahli (Project Manager & Pipeline Engineer)",
                "dokumen_diperlukan": "CV, SKA Madya Migas, Ijazah",
                "parameter_evaluasi": "Pengalaman minimal 5 tahun di bidang konstruksi perpipaan migas",
                "bobot": 30
            },
            {
                "no": 3,
                "kriteria": "Peralatan Utama (Welding Machine, Excavator, Crane)",
                "dokumen_diperlukan": "Daftar Kepemilikan/Sewa & Sertifikat SILO Migas",
                "parameter_evaluasi": "Ketersediaan dan kelaikan peralatan utama",
                "bobot": 20
            },
            {
                "no": 4,
                "kriteria": "Manajemen Mutu & Quality Plan",
                "dokumen_diperlukan": "Manual Mutu ISO 9001 & Inspection Test Plan (ITP)",
                "parameter_evaluasi": "Kelengkapan prosedur NDT dan pengujian mutu",
                "bobot": 20
            }
        ],
        "rancangan_kontrak": {
            "judul_pekerjaan": "PENGADAAN DAN PEMASANGAN PIPA TRANSMISI MINYAK DAN GAS REGIONAL KALIMANTAN",
            "nomor_kontrak": "KTR-042/PND970000/2026-S7",
            "durasi_pekerjaan_hari": 180,
            "masa_pemeliharaan_hari": 90,
            "sistem_pembayaran": "Milestone Progress bulanan berdasarkan Berita Acara Progres Fisik yang disetujui Pengawas Pekerjaan.",
            "denda_per_mil": 1.0,
            "lingkup_pekerjaan": [
                "Penyediaan material pipa baja dan fitting sesuai API 5L spesifikasi teknis Pertamina",
                "Pekerjaan sipil galian tanah, penggelaran (stringing), dan penimbunan kembali (backfilling)",
                "Pekerjaan pengelasan pipa (welding) dan pengujian tanpa merusak (Non-Destructive Testing / NDT)",
                "Pembersihan pipa (swabbing), pengeringan, dan uji tekan hidrostatik (Hydrostatic Test)",
                "Pekerjaan perlindungan korosi (Coating dan Cathodic Protection system)"
            ]
        }
    }

    print("Rendering Dokumen Tender to DOCX...")
    output_path = render_tender_to_docx(sample_data, "Test_Dokumen_Tender.docx")
    print(f"Success! Generated file at: {output_path}")

    # Inspect generated docx
    doc = Document(output_path)
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    print(f"Total Sections: {len(doc.sections)}")

    # Verify key text exists in paragraphs and tables
    full_text = "\n".join(p.text for p in doc.paragraphs)
    table_text = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    all_doc_text = full_text + " " + table_text

    assert "DOKUMEN TENDER" in all_doc_text, "Missing DOKUMEN TENDER title"
    assert "BAGIAN A" in all_doc_text, "Missing BAGIAN A"
    assert "BAB I KETENTUAN KHUSUS IKPP" in all_doc_text, "Missing BAB I"
    assert "BAB II KETENTUAN UMUM IKPP" in all_doc_text, "Missing BAB II"
    assert "LAMPIRAN IKPP" in all_doc_text, "Missing LAMPIRAN IKPP"
    assert "BAGIAN B" in all_doc_text, "Missing BAGIAN B"
    assert "PERJANJIAN PELAKSANAAN PEKERJAAN" in all_doc_text, "Missing Contract Agreement"
    assert "PAKTA INTEGRITAS PELAKSANA KONTRAK" in all_doc_text, "Missing PI-07"
    assert "SP3MK" in all_doc_text, "Missing SP3MK"
    print("All assertions passed successfully!")

if __name__ == "__main__":
    test_tender_rendering()
