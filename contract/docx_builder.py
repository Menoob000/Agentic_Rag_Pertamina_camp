"""
contract/docx_builder.py — DOCX Builder for Rancangan Kontrak & Amendemen
==========================================================================
Membangun dokumen Word (.docx) dari JSON draft yang di-generate oleh LLM.
Menggunakan library python-docx dan meminjam styling XML dari tender/docx.py.
"""

import os
from pathlib import Path
from typing import Dict, Any

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ──────────────────────────────────────────────────────────────
# Setup & Constants
# ──────────────────────────────────────────────────────────────

OUTPUT_DIR = Path("./output")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# Color Palette (Pertamina Corporate theme)
COLOR_NAVY = "002D62"
COLOR_HEADER_BG = "D9E2F3"
COLOR_ALT_BG = "F7F9FC"
COLOR_WHITE = "FFFFFF"

# ──────────────────────────────────────────────────────────────
# XML Formatting Helpers (Reused from tender/docx.py)
# ──────────────────────────────────────────────────────────────

def _set_cell_shading(cell, color_hex: str):
    """Set background fill color for a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def _set_cell_border(cell, **kwargs):
    """Set cell borders: top, bottom, left, right dicts with {val, sz, color}."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for edge in ("top", "bottom", "left", "right"):
        if edge in kwargs:
            attrs = kwargs[edge]
            element = parse_xml(
                f'<w:{edge} {nsdecls("w")} w:val="{attrs.get("val", "single")}" '
                f'w:sz="{attrs.get("sz", "4")}" w:space="0" '
                f'w:color="{attrs.get("color", "000000")}"/>'
            )
        else:
            element = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
        tcBorders.append(element)
    tcPr.append(tcBorders)

def add_disclaimer_header(doc: Document, text: str = "DRAFT DOKUMEN (AI-GENERATED) – HARAP DIREVIEW OLEH TIM LEGAL"):
    """Adds a bold red disclaimer at the top of the document."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.font.size = Pt(11)


# ──────────────────────────────────────────────────────────────
# Builder 1: Rancangan Kontrak Definitif
# ──────────────────────────────────────────────────────────────

def render_kontrak_to_docx(draft: Dict[str, Any], output_filename: str) -> str:
    """Render KontrakDefinitif JSON to a DOCX file."""
    doc = Document()
    
    # Base styling
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    add_disclaimer_header(doc)
    
    # ── 1. Cover / Header ──
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("RANCANGAN KONTRAK\nPERJANJIAN PEMBORONGAN PEKERJAAN\n\n")
    run.bold = True
    run.font.size = Pt(14)
    
    judul = doc.add_paragraph()
    judul.alignment = WD_ALIGN_PARAGRAPH.CENTER
    judul.add_run(draft.get("judul_pekerjaan", "PEKERJAAN")).bold = True
    
    doc.add_paragraph(f"Nomor Kontrak : {draft.get('nomor_kontrak', '-')}", style='Normal')
    doc.add_paragraph(f"Tanggal       : {draft.get('tanggal_kontrak', '-')}", style='Normal')
    doc.add_paragraph(f"Jenis Kontrak : {draft.get('jenis_kontrak', '-')}", style='Normal')
    
    doc.add_page_break()
    
    # ── 2. Identitas Para Pihak ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Pada hari ini, tanggal ").bold = False
    p.add_run(draft.get('tanggal_kontrak', '-')).bold = True
    p.add_run(", bertempat di ").bold = False
    p.add_run(draft.get('lokasi_pekerjaan', '-')).bold = True
    p.add_run(", kami yang bertanda tangan di bawah ini:\n").bold = False
    
    pertamina = draft.get("pihak_pertamina", {})
    p.add_run(f"\n1. {pertamina.get('nama_perusahaan', 'PT PERTAMINA PATRA NIAGA')}").bold = True
    p.add_run(f", berkedudukan di {pertamina.get('alamat', '-')}, dalam hal ini diwakili oleh ")
    p.add_run(f"{pertamina.get('diwakili_oleh', '-')}").bold = True
    p.add_run(f" selaku {pertamina.get('jabatan', '-')} (selanjutnya disebut \"PERUSAHAAN\").\n")
    
    vendor = draft.get("pihak_vendor", {})
    p.add_run(f"\n2. {vendor.get('nama_perusahaan', '-')}").bold = True
    p.add_run(f", suatu {vendor.get('bentuk_badan_usaha', '-')}, berkedudukan di {vendor.get('alamat', '-')}, ")
    p.add_run(f"dalam hal ini diwakili oleh {vendor.get('diwakili_oleh', '-')} selaku {vendor.get('jabatan', '-')} ")
    p.add_run("(selanjutnya disebut \"PENYEDIA\").\n")
    
    p.add_run("\nPERUSAHAAN dan PENYEDIA secara bersama-sama disebut PARA PIHAK dan secara sendiri-sendiri disebut PIHAK. ")
    p.add_run("PARA PIHAK sepakat untuk mengikatkan diri dalam Perjanjian ini dengan syarat dan ketentuan sebagai berikut:")
    
    doc.add_page_break()
    
    # ── 3. Pasal-Pasal ──
    for pasal in draft.get("pasal", []):
        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.add_run(f"PASAL {pasal.get('nomor', '')}\n").bold = True
        h.add_run(f"{pasal.get('judul', '').upper()}").bold = True
        
        for ayat in pasal.get("ayat", []):
            a = doc.add_paragraph(str(ayat))
            a.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            a.paragraph_format.space_after = Pt(6)
            
    doc.add_page_break()
    
    # ── 4. Penandatangan (Tanda Tangan) ──
    sig_p = doc.add_paragraph("Demikian Kontrak ini dibuat dalam rangkap 2 (dua), bermeterai cukup dan memiliki kekuatan hukum yang sama.")
    sig_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    
    # Remove borders for signature table
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(cell, top={"val": "none"}, bottom={"val": "none"}, left={"val": "none"}, right={"val": "none"})
            
    cells_r0 = table.rows[0].cells
    cells_r0[0].text = "PERUSAHAAN\nPT PERTAMINA PATRA NIAGA"
    cells_r0[1].text = f"PENYEDIA\n{vendor.get('nama_perusahaan', '-')}"
    for c in cells_r0:
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].runs[0].bold = True
        
    cells_r1 = table.rows[1].cells
    cells_r1[0].text = "\n\n\n\n" # Space for signature
    
    cells_r3 = table.rows[3].cells
    cells_r3[0].text = f"{pertamina.get('diwakili_oleh', '________________')}\n{pertamina.get('jabatan', '-')}"
    cells_r3[1].text = f"{vendor.get('diwakili_oleh', '________________')}\n{vendor.get('jabatan', '-')}"
    for c in cells_r3:
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].runs[0].underline = True

    # Save
    if not output_filename.endswith(".docx"):
        output_filename += ".docx"
    output_path = OUTPUT_DIR / output_filename
    doc.save(output_path)
    return str(output_path)


# ──────────────────────────────────────────────────────────────
# Builder 2: Draft Amendemen / Addendum
# ──────────────────────────────────────────────────────────────

def render_amendemen_to_docx(draft: Dict[str, Any], output_filename: str) -> str:
    """Render DraftAmendemen JSON to a DOCX file."""
    doc = Document()
    
    # Base styling
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    add_disclaimer_header(doc)
    
    # ── 1. Cover / Header ──
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"DRAFT {draft.get('nomor_addendum', 'ADDENDUM').upper()}\n")
    run.bold = True
    run.font.size = Pt(14)
    
    asal = draft.get("kontrak_asal", {})
    judul = doc.add_paragraph()
    judul.alignment = WD_ALIGN_PARAGRAPH.CENTER
    judul.add_run(f"ATAS KONTRAK NOMOR: {asal.get('nomor_kontrak', '-')}\n").bold = True
    judul.add_run(f"PEKERJAAN: {asal.get('judul_pekerjaan', '-').upper()}").bold = True
    
    doc.add_paragraph(f"Tanggal Berlaku Efektif: {draft.get('tanggal_berlaku_efektif', '-')}", style='Normal')
    
    doc.add_page_break()
    
    # ── 2. Konsideran ──
    k_head = doc.add_paragraph()
    k_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    k_head.add_run("KONSIDERAN").bold = True
    
    konsideran = draft.get("konsideran", {})
    
    def add_konsideran_section(title_text, items):
        if not items: return
        p = doc.add_paragraph()
        p.add_run(title_text).bold = True
        for idx, item in enumerate(items, 1):
            p = doc.add_paragraph(f"{idx}. {item}")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(1.25)
            
    add_konsideran_section("MENGINGAT:", konsideran.get("mengingat", []))
    add_konsideran_section("MEMPERHATIKAN:", konsideran.get("memperhatikan", []))
    add_konsideran_section("MENIMBANG:", konsideran.get("menimbang", []))
    
    doc.add_page_break()
    
    # ── 3. Detail Perubahan ──
    c_head = doc.add_paragraph()
    c_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c_head.add_run("PERUBAHAN").bold = True
    
    p_intro = doc.add_paragraph("PARA PIHAK sepakat untuk melakukan perubahan terhadap Kontrak Asal sebagai berikut:")
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    perubahan_list = draft.get("perubahan", [])
    for idx, item in enumerate(perubahan_list, 1):
        pasal_ref = item.get("pasal_referensi", f"Perubahan {idx}")
        p = doc.add_paragraph(f"{idx}. Mengubah {pasal_ref} sehingga berbunyi sebagai berikut:")
        p.runs[0].bold = True
        p.paragraph_format.space_after = Pt(6)
        
        # Tabel Lama vs Baru
        table = doc.add_table(rows=2, cols=2)
        table.style = 'Table Grid'
        
        # Header Tabel
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Semula Berbunyi (Kontrak Asal)"
        hdr_cells[1].text = "Diubah Menjadi (Addendum)"
        for c in hdr_cells:
            c.paragraphs[0].runs[0].bold = True
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_cell_shading(c, COLOR_HEADER_BG)
            
        # Isi Tabel
        row_cells = table.rows[1].cells
        row_cells[0].text = item.get("isi_lama", "-")
        row_cells[1].text = item.get("isi_baru", "-")
        
        # Alasan
        doc.add_paragraph(f"Dasar/Alasan: {item.get('alasan_perubahan', '-')}\n").runs[0].italic = True

    # ── 4. Pasal Penutup ──
    doc.add_paragraph(draft.get("pasal_penutup", "Ketentuan-ketentuan lain dalam Kontrak tetap berlaku."))
    
    # ── 5. Penandatangan ──
    doc.add_paragraph("\n")
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(cell, top={"val": "none"}, bottom={"val": "none"}, left={"val": "none"}, right={"val": "none"})
            
    cells_r0 = table.rows[0].cells
    cells_r0[0].text = "PERUSAHAAN\nPT PERTAMINA PATRA NIAGA"
    cells_r0[1].text = f"PENYEDIA\n{asal.get('nama_vendor', '-')}"
    for c in cells_r0:
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].runs[0].bold = True
        
    cells_r1 = table.rows[1].cells
    cells_r1[0].text = "\n\n\n\n"
    
    cells_r3 = table.rows[3].cells
    cells_r3[0].text = f"{draft.get('ttd_pertamina_nama', '________________')}\n{draft.get('ttd_pertamina_jabatan', '-')}"
    cells_r3[1].text = f"{draft.get('ttd_vendor_nama', '________________')}\n{draft.get('ttd_vendor_jabatan', 'Direktur Utama')}"
    for c in cells_r3:
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].runs[0].underline = True

    # Save
    if not output_filename.endswith(".docx"):
        output_filename += ".docx"
    output_path = OUTPUT_DIR / output_filename
    doc.save(output_path)
    return str(output_path)
