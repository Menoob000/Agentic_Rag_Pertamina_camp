"""
RKS DOCX Renderer
Converts intermediate RKS format (JSON/dict) into a professionally formatted
.docx file with formal Pertamina document elements:
- Cover page with page border, project title, CSMS risk checkbox, logo
- Approval/sign-off sheet
- Table of contents placeholder
- Content pages with header/footer
"""

from docx.text.paragraph import Paragraph
import json
import os
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# ──────────────────────────────────────────────────────────────
# Configuration
# ──────────────────────────────────────────────────────────────

ASSETS_DIR = Path("./assets")
OUTPUT_DIR = Path("./output")
CONFIG_PATH = ASSETS_DIR / "doc_config.json"
LOGO_PATH = ASSETS_DIR / "logo_pertamina.png"


def _load_config() -> dict:
    """Load document configuration from assets/doc_config.json."""
    if CONFIG_PATH.exists():
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    # Fallback defaults
    return {
        "perusahaan": {"nama_default": "PT PERTAMINA (PERSERO)"},
        "styling": {
            "font_body": "Arial",
            "font_heading": "Arial",
            "font_size_body": 12,
            "font_size_heading": 14,
            "margin_top_cm": 3.0,
            "margin_bottom_cm": 2.5,
            "margin_left_cm": 3.0,
            "margin_right_cm": 2.5,
            "line_spacing": 1.5,
        },
        "elemen_formal": {
            "halaman_sampul": True,
            "lembar_pengesahan": True,
            "daftar_isi": True,
            "page_border": True,
            "header_logo": True,
            "footer_nomor_dok": True,
        },
    }


# ──────────────────────────────────────────────────────────────
# Helper Functions
# ──────────────────────────────────────────────────────────────

def _set_cell_shading(cell, color: str):
    """Set background color for a table cell."""
    shading_elm = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>'
    )
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _set_cell_border(cell, **kwargs):
    """Set borders on a table cell. kwargs: top, bottom, left, right, each a dict with sz, val, color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')

    for edge, attrs in kwargs.items():
        element = parse_xml(
            f'<w:{edge} {nsdecls("w")} w:val="{attrs.get("val", "single")}" '
            f'w:sz="{attrs.get("sz", "4")}" w:space="0" '
            f'w:color="{attrs.get("color", "000000")}"/>'
        )
        tcBorders.append(element)

    tcPr.append(tcBorders)


def _add_page_border(section):
    """Add a page border to the section."""
    sectPr = section._sectPr
    pgBorders = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'  <w:top w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        f'</w:pgBorders>'
    )
    sectPr.append(pgBorders)


def _remove_page_border(section):
    """Remove page border from the section."""
    sectPr = section._sectPr
    pgBorders = sectPr.find(qn('w:pgBorders'))
    if pgBorders is not None:
        sectPr.remove(pgBorders)


def _set_paragraph_font(paragraph, font_name: str, font_size: int, bold: bool = False, color: Optional[RGBColor] = None):
    """Apply font formatting to all runs in a paragraph."""
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        # Ensure East Asian font is set too
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)


def _add_formatted_paragraph(doc, text: str, font_name: str = "Arial", font_size: int = 12,
                              bold: bool = False, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                              space_before: int = 0, space_after: int = 6) -> 'Paragraph':
    """Add a paragraph with formatting."""
    para = doc.add_paragraph()
    para.alignment = alignment
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)
    run = para.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    return para


# ──────────────────────────────────────────────────────────────
# Cover Page
# ──────────────────────────────────────────────────────────────

def _render_cover_page(doc: Document, draft_data: dict, config: dict):
    """Render the RKS cover page based on the Pertamina RKS standard."""
    cover = draft_data.get("cover_page") or {}
    styling = config.get("styling") or {}
    font = styling.get("font_heading", "Arial")

    section = doc.sections[0]

    # Add page border if enabled
    if (config.get("elemen_formal") or {}).get("page_border", True):
        _add_page_border(section)

    # ── Title: "RENCANA KERJA DAN SYARAT-SYARAT (RKS)" in a bordered box ──
    doc.add_paragraph()  # Top spacing

    # Create a 1x1 table for the title box
    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_cell = title_table.cell(0, 0)

    # Title text
    title_para = title_cell.paragraphs[0]
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("RENCANA KERJA DAN SYARAT-SYARAT\n(RKS)")
    title_run.font.name = font
    title_run.font.size = Pt(16)
    title_run.font.bold = True

    # Set cell borders
    _set_cell_border(
        title_cell,
        top={"sz": "8", "val": "single", "color": "000000"},
        bottom={"sz": "8", "val": "single", "color": "000000"},
        left={"sz": "8", "val": "single", "color": "000000"},
        right={"sz": "8", "val": "single", "color": "000000"},
    )

    doc.add_paragraph()  # Spacing

    # ── Project description in a bordered box ──
    judul_pekerjaan = cover.get("judul_pekerjaan", "DESKRIPSI PEKERJAAN")

    desc_table = doc.add_table(rows=1, cols=1)
    desc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    desc_cell = desc_table.cell(0, 0)

    desc_para = desc_cell.paragraphs[0]
    desc_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc_run = desc_para.add_run(judul_pekerjaan)
    desc_run.font.name = font
    desc_run.font.size = Pt(12)
    desc_run.font.bold = True

    _set_cell_border(
        desc_cell,
        top={"sz": "8", "val": "single", "color": "000000"},
        bottom={"sz": "8", "val": "single", "color": "000000"},
        left={"sz": "8", "val": "single", "color": "000000"},
        right={"sz": "8", "val": "single", "color": "000000"},
    )

    doc.add_paragraph()  # Spacing

    # ── Bidder List No ──
    bidder_para = _add_formatted_paragraph(
        doc,
        f"Bidder List No: {'_' * 30}",
        font_name=font, font_size=12, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=12, space_after=12
    )

    # ── RESIKO PEKERJAAN (ASPEK CSMS) ──
    _add_formatted_paragraph(
        doc,
        "RESIKO PEKERJAAN (ASPEK CSMS):",
        font_name=font, font_size=12, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=12, space_after=6
    )

    # Checkboxes
    resiko = cover.get("resiko_csms", "HIGH").upper()
    check_high = "☑" if resiko == "HIGH" else "☐"
    check_mid = "☑" if resiko == "MIDDLE" else "☐"
    check_low = "☑" if resiko == "LOW" else "☐"

    _add_formatted_paragraph(
        doc,
        f"{check_high} HIGH     {check_mid} MIDDLE     {check_low} LOW",
        font_name=font, font_size=12, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0, space_after=18
    )

    # ── Logo ──
    if LOGO_PATH.exists():
        logo_para = doc.add_paragraph()
        logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = logo_para.add_run()
        run.add_picture(str(LOGO_PATH), width=Cm(5))
    else:
        # Placeholder text if logo not found
        _add_formatted_paragraph(
            doc,
            "PERTAMINA",
            font_name=font, font_size=24, bold=True,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            space_before=24, space_after=24
        )

    # ── Spacing to push company name toward bottom ──
    for _ in range(4):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(12)
        spacer.paragraph_format.space_after = Pt(12)

    # ── Company name at bottom ──
    nama_perusahaan = cover.get(
        "nama_perusahaan",
        (config.get("perusahaan") or {}).get("nama_default", "PT PERTAMINA (PERSERO)")
    )
    _add_formatted_paragraph(
        doc,
        nama_perusahaan,
        font_name=font, font_size=12, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0, space_after=0
    )

    # Section break after cover to isolate page borders
    new_section = doc.add_section()
    _remove_page_border(new_section)


# ──────────────────────────────────────────────────────────────
# Approval Sheet (Lembar Pengesahan)
# ──────────────────────────────────────────────────────────────

def _render_approval_sheet(doc: Document, draft_data: dict, config: dict):
    """Render the approval/sign-off sheet."""
    pengesahan = draft_data.get("pengesahan") or {}
    font = (config.get("styling") or {}).get("font_heading", "Arial")

    _add_formatted_paragraph(
        doc, "LEMBAR PENGESAHAN",
        font_name=font, font_size=14, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=24, space_after=24
    )

    # Create approval table
    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    headers = ["", "Nama / Jabatan", "Tanda Tangan"]
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.name = font
            run.font.size = Pt(11)
        _set_cell_shading(cell, "D9E2F3")

    # Data rows
    roles = [
        ("Disusun oleh", pengesahan.get("disusun_oleh") or {}),
        ("Diperiksa oleh", pengesahan.get("diperiksa_oleh") or {}),
        ("Disetujui oleh", pengesahan.get("disetujui_oleh") or {}),
    ]

    for row_idx, (role, data) in enumerate(roles, start=1):
        # Role column
        cell_role = table.cell(row_idx, 0)
        cell_role.text = role
        cell_role.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell_role.paragraphs[0].runs:
            run.font.name = font
            run.font.size = Pt(11)
            run.font.bold = True

        # Name/Position column
        cell_name = table.cell(row_idx, 1)
        nama = data.get("nama", "") or ""
        jabatan = data.get("jabatan", "") or ""
        name_text = f"\n\n\n{nama}\n{jabatan}" if nama else f"\n\n\n{'_' * 20}\n({'_' * 20})"
        cell_name.text = name_text
        cell_name.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell_name.paragraphs[0].runs:
            run.font.name = font
            run.font.size = Pt(10)

        # Signature column
        cell_sign = table.cell(row_idx, 2)
        cell_sign.text = f"\n\n\n{'_' * 20}\nTanggal: ___/___/______"
        cell_sign.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell_sign.paragraphs[0].runs:
            run.font.name = font
            run.font.size = Pt(10)

    # Apply borders to all cells
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(
                cell,
                top={"sz": "4", "val": "single", "color": "000000"},
                bottom={"sz": "4", "val": "single", "color": "000000"},
                left={"sz": "4", "val": "single", "color": "000000"},
                right={"sz": "4", "val": "single", "color": "000000"},
            )

    doc.add_page_break()


# ──────────────────────────────────────────────────────────────
# Table of Contents Placeholder
# ──────────────────────────────────────────────────────────────

def _render_toc(doc: Document, config: dict):
    """Add a Table of Contents placeholder (user updates in Word with F9)."""
    font = (config.get("styling") or {}).get("font_heading", "Arial")

    _add_formatted_paragraph(
        doc, "DAFTAR ISI",
        font_name=font, font_size=14, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=24, space_after=18
    )

    # Insert TOC field
    para = doc.add_paragraph()
    run = para.add_run()
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._element.append(fldChar1)

    run2 = para.add_run()
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>')
    run2._element.append(instrText)

    run3 = para.add_run()
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    run3._element.append(fldChar2)

    run4 = para.add_run("(Tekan F9 di Microsoft Word untuk update Daftar Isi)")
    run4.font.name = font
    run4.font.size = Pt(11)
    run4.font.color.rgb = RGBColor(128, 128, 128)

    run5 = para.add_run()
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run5._element.append(fldChar3)

    doc.add_page_break()


# ──────────────────────────────────────────────────────────────
# Header & Footer
# ──────────────────────────────────────────────────────────────

def _setup_header_footer(section, draft_data: dict, config: dict):
    """Set up header with logo/company name and footer with doc number/page."""
    meta = draft_data.get("metadata") or {}
    company = (draft_data.get("cover_page") or {}).get(
        "nama_perusahaan",
        (config.get("perusahaan") or {}).get("nama_default", "PT PERTAMINA (PERSERO)")
    )
    font = (config.get("styling") or {}).get("font_heading", "Arial")
    nomor_dok = meta.get("nomor_dokumen", "")

    # ── Header ──
    header = section.header
    header.is_linked_to_previous = False

    header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Add logo if exists
    if LOGO_PATH.exists():
        run = header_para.add_run()
        run.add_picture(str(LOGO_PATH), height=Cm(1.2))
        header_para.add_run("  ")

    run = header_para.add_run(company)
    run.font.name = font
    run.font.size = Pt(9)
    run.font.bold = True

    # Add bottom border to header
    pPr = header_para._element.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

    # ── Footer ──
    footer = section.footer
    footer.is_linked_to_previous = False

    footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if nomor_dok:
        run = footer_para.add_run(f"No. Dok: {nomor_dok}  |  ")
        run.font.name = font
        run.font.size = Pt(8)

    # Page number field
    run = footer_para.add_run("Hal. ")
    run.font.name = font
    run.font.size = Pt(8)

    # Current page number
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run_field = footer_para.add_run()
    run_field._element.append(fldChar1)

    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    run_field2 = footer_para.add_run()
    run_field2._element.append(instrText)

    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run_field3 = footer_para.add_run()
    run_field3._element.append(fldChar2)

    # Add top border to footer
    pPr = footer_para._element.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="1" w:color="000000"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


# ──────────────────────────────────────────────────────────────
# Content Rendering
# ──────────────────────────────────────────────────────────────

def _render_content(doc: Document, draft_data: dict, config: dict):
    """Render all bab and sub-bab content."""
    styling = config.get("styling") or {}
    font_body = styling.get("font_body", "Arial")
    font_heading = styling.get("font_heading", "Arial")
    size_body = styling.get("font_size_body", 12)
    size_heading = styling.get("font_size_heading", 14)
    line_spacing = styling.get("line_spacing", 1.5)

    for bab in draft_data.get("bab") or []:
        nomor = bab.get("nomor", "")
        judul = bab.get("judul", "")

        # Bab heading (Heading 1)
        heading = doc.add_heading(f"BAB {nomor} - {judul}", level=1)
        for run in heading.runs:
            run.font.name = font_heading
            run.font.size = Pt(size_heading)
            run.font.color.rgb = RGBColor(0, 0, 0)

        for sub in bab.get("sub_bab") or []:
            sub_nomor = sub.get("nomor", "")
            sub_judul = sub.get("judul", "")
            tipe = sub.get("tipe", "paragraf")
            konten = sub.get("konten", "")

            # Sub-bab heading (Heading 2)
            sub_heading = doc.add_heading(f"{sub_nomor} {sub_judul}", level=2)
            for run in sub_heading.runs:
                run.font.name = font_heading
                run.font.size = Pt(size_body)
                run.font.color.rgb = RGBColor(0, 0, 0)

            # Render content based on type
            if tipe == "heading_only" or not konten:
                continue
            elif tipe == "paragraf":
                _render_paragraph(doc, konten, font_body, size_body, line_spacing)
            elif tipe == "numbered_list":
                _render_list(doc, konten, font_body, size_body, numbered=True)
            elif tipe == "bullet_list":
                _render_list(doc, konten, font_body, size_body, numbered=False)
            elif tipe == "table":
                _render_table(doc, konten, font_body)
            else:
                # Default to paragraph
                _render_paragraph(doc, str(konten), font_body, size_body, line_spacing)


def _render_paragraph(doc: Document, text, font_name: str, font_size: int, line_spacing: float):
    """Render paragraph content."""
    if isinstance(text, list):
        text = "\n".join(text)
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = line_spacing
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(str(text))
    run.font.name = font_name
    run.font.size = Pt(font_size)


def _render_list(doc: Document, items, font_name: str, font_size: int, numbered: bool = True):
    """Render a numbered or bullet list."""
    if isinstance(items, str):
        items = [item.strip() for item in items.split("\n") if item.strip()]
    elif not isinstance(items, list):
        items = [str(items)]

    style = 'List Number' if numbered else 'List Bullet'
    for item in items:
        para = doc.add_paragraph(str(item), style=style)
        for run in para.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)


def _render_table(doc: Document, table_data, font_name: str):
    """Render a table from dict with 'headers' and 'rows'."""
    if isinstance(table_data, str):
        # If it's a string, render as paragraph instead
        para = doc.add_paragraph(table_data)
        for run in para.runs:
            run.font.name = font_name
            run.font.size = Pt(10)
        return

    if not isinstance(table_data, dict):
        return

    headers = table_data.get("headers") or []
    rows = table_data.get("rows") or []

    if not headers:
        return

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = str(header)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.name = font_name
            run.font.size = Pt(10)
        _set_cell_shading(cell, "D9E2F3")

    # Data rows
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_value in enumerate(row_data):
            if col_idx < len(headers):
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = str(cell_value)
                for run in cell.paragraphs[0].runs:
                    run.font.name = font_name
                    run.font.size = Pt(10)

    # Apply borders
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(
                cell,
                top={"sz": "4", "val": "single", "color": "000000"},
                bottom={"sz": "4", "val": "single", "color": "000000"},
                left={"sz": "4", "val": "single", "color": "000000"},
                right={"sz": "4", "val": "single", "color": "000000"},
            )

    doc.add_paragraph()  # Spacing after table


# ──────────────────────────────────────────────────────────────
# Main Render Function
# ──────────────────────────────────────────────────────────────

def render_to_docx(draft_data: dict, output_filename: str = "") -> str:
    """
    Render the complete RKS document to a .docx file.
    
    Args:
        draft_data: Complete RKS data in intermediate format
        output_filename: Optional filename (auto-generated if empty)
    
    Returns:
        str: Absolute path to the generated .docx file
    """
    config = _load_config()
    styling = config.get("styling") or {}
    elemen = config.get("elemen_formal") or {}

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # ── Set document-wide margins ──
    for section in doc.sections:
        section.top_margin = Cm(styling.get("margin_top_cm", 3.0))
        section.bottom_margin = Cm(styling.get("margin_bottom_cm", 2.5))
        section.left_margin = Cm(styling.get("margin_left_cm", 3.0))
        section.right_margin = Cm(styling.get("margin_right_cm", 2.5))

    # ── Set default font styles ──
    style = doc.styles['Normal']
    font = style.font
    font.name = styling.get("font_body", "Arial")
    font.size = Pt(styling.get("font_size_body", 12))

    # ── Render sections ──

    # 1. Cover page
    if elemen.get("halaman_sampul", True):
        _render_cover_page(doc, draft_data, config)

    # 2. Approval sheet
    if elemen.get("lembar_pengesahan", True):
        _render_approval_sheet(doc, draft_data, config)

    # 3. Table of contents
    if elemen.get("daftar_isi", True):
        _render_toc(doc, config)

    # 4. Setup header/footer for content pages
    # Add a new section for content (to separate from front matter)
    # Ensure there is no page border
    new_section = doc.add_section()
    new_section.top_margin = Cm(styling.get("margin_top_cm", 3.0))
    new_section.bottom_margin = Cm(styling.get("margin_bottom_cm", 2.5))
    new_section.left_margin = Cm(styling.get("margin_left_cm", 3.0))
    new_section.right_margin = Cm(styling.get("margin_right_cm", 2.5))
    _remove_page_border(new_section)

    if elemen.get("header_logo", True) or elemen.get("footer_nomor_dok", True):
        _setup_header_footer(new_section, draft_data, config)

    # 5. Content (bab & sub-bab)
    _render_content(doc, draft_data, config)

    # ── Save ──
    if not output_filename:
        cover = draft_data.get("cover_page") or {}
        name_part = cover.get("judul_pekerjaan", "RKS_Document")
        safe_name = "".join(
            c if c.isalnum() or c in (" ", "-", "_") else "_"
            for c in name_part
        )[:60].strip()
        output_filename = f"RKS_{safe_name}.docx"

    if not output_filename.endswith(".docx"):
        output_filename += ".docx"

    output_path = OUTPUT_DIR / output_filename
    doc.save(str(output_path))

    return str(output_path.resolve())
