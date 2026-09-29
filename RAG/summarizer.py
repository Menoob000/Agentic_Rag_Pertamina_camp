import json
import os
import tempfile
import io
from datetime import datetime, date, timedelta
from typing import Optional, Dict, Any, List

from langchain_core.messages import SystemMessage, HumanMessage
from langchain_openrouter import ChatOpenRouter
from dotenv import load_dotenv

load_dotenv()

_doc_converter = None

INDONESIAN_DAYS = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
INDONESIAN_MONTHS = [
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember"
]


def format_id_date(d: date, with_day: bool = True) -> str:
    """Format python date to Indonesian standard string."""
    day_name = INDONESIAN_DAYS[d.weekday()]
    month_name = INDONESIAN_MONTHS[d.month - 1]
    if with_day:
        return f"{day_name}, {d.day:02d} {month_name} {d.year}"
    return f"{d.day:02d} {month_name} {d.year}"


def parse_date_safely(date_input: Any) -> date:
    """Safely parse various date representations into a python date object."""
    if isinstance(date_input, datetime):
        return date_input.date()
    if isinstance(date_input, date):
        return date_input

    str_val = str(date_input or "").strip()
    if not str_val:
        return date.today()

    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(str_val, fmt).date()
        except ValueError:
            pass

    try:
        parts = str_val.replace(",", "").split()
        day = None
        month = None
        year = None
        for p in parts:
            if p.isdigit():
                val = int(p)
                if val > 1900:
                    year = val
                elif day is None and 1 <= val <= 31:
                    day = val
            else:
                for idx, m in enumerate(INDONESIAN_MONTHS):
                    if m.lower() in p.lower():
                        month = idx + 1
                        break
        if day and month and year:
            return date(year, month, day)
    except Exception:
        pass

    return date.today()


def calculate_tender_schedules(base_d: date) -> Dict[str, str]:
    """Calculate Pre-bid and bid submission dates from base tender date."""
    # Pre-bid meeting: base_date + 4 days, avoid weekend
    prebid_d = base_d + timedelta(days=4)
    if prebid_d.weekday() == 5:  # Saturday
        prebid_d += timedelta(days=2)
    elif prebid_d.weekday() == 6:  # Sunday
        prebid_d += timedelta(days=1)

    pemasukan_mulai = prebid_d
    pemasukan_selesai = prebid_d + timedelta(days=7)
    if pemasukan_selesai.weekday() == 5:
        pemasukan_selesai += timedelta(days=2)
    elif pemasukan_selesai.weekday() == 6:
        pemasukan_selesai += timedelta(days=1)

    return {
        "prebid_tanggal": format_id_date(prebid_d, with_day=True),
        "prebid_waktu": "10.00 WITA",
        "prebid_tempat": "Microsoft Teams Meeting dengan link yang disampaikan melalui email Undangan Prebid Meeting",
        "pemasukan_mulai": format_id_date(pemasukan_mulai, with_day=True),
        "pemasukan_selesai": format_id_date(pemasukan_selesai, with_day=True),
    }


def get_converter():
    global _doc_converter
    if _doc_converter is None:
        from docling.document_converter import DocumentConverter
        _doc_converter = DocumentConverter()
    return _doc_converter


def extract_text_from_pdfs(pdf_bytes_list: list[bytes]) -> str:
    """Extract text from a list of PDF byte arrays using pypdf (ultra-fast) with Docling fallback for scanned PDFs."""
    combined_text = []

    for pdf_bytes in pdf_bytes_list:
        # 1. Fast extraction with pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
            pages_text = []
            for p in reader.pages:
                pt = p.extract_text()
                if pt:
                    pages_text.append(pt)
            if pages_text:
                extracted = "\n\n".join(pages_text).strip()
                if len(extracted) > 100:
                    print(f"  [pypdf] Berhasil mengekstrak {len(reader.pages)} halaman ({len(extracted)} karakter) secara instan.")
                    combined_text.append(extracted)
                    continue
        except Exception as e:
            print(f"  [pypdf] Melewati pypdf: {e}")

        # 2. Fallback to Docling if text could not be extracted (e.g. scanned image PDF)
        converter = get_converter()
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_file:
            tmp_file.write(pdf_bytes)
            tmp_file_path = tmp_file.name

        try:
            print("  [Docling OCR] Mulai memproses PDF terpindai...")
            result = converter.convert(tmp_file_path)
            markdown_text = result.document.export_to_markdown()
            if markdown_text:
                combined_text.append(markdown_text)
            print("  [Docling OCR] Selesai memproses dokumen.")
        except Exception as e:
            print(f"  [Docling OCR] Error saat membaca PDF: {e}")
        finally:
            if os.path.exists(tmp_file_path):
                os.remove(tmp_file_path)

    return "\n\n---\n\n".join(combined_text)


def extract_text_from_excel(excel_bytes: bytes) -> str:
    """Extract tabular text from Excel (.xlsx) using openpyxl."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
        sheets_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_text = []
            for row in ws.iter_rows(values_only=True):
                if any(v is not None and str(v).strip() != "" for v in row):
                    formatted_row = [str(v).strip() if v is not None else "" for v in row]
                    while formatted_row and not formatted_row[-1]:
                        formatted_row.pop()
                    if formatted_row:
                        rows_text.append(" | ".join(formatted_row))
            if rows_text:
                sheet_content = f"### Sheet: {sheet_name}\n" + "\n".join(rows_text)
                sheets_text.append(sheet_content)
        return "\n\n".join(sheets_text)
    except Exception as e:
        print(f"  [Excel Extractor] Error saat membaca file Excel: {e}")
        return ""


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extract text from PDF or Excel file based on extension."""
    if not file_bytes:
        return ""
    fname = filename.lower()
    if fname.endswith(".pdf"):
        return extract_text_from_pdfs([file_bytes])
    elif fname.endswith(".xlsx") or fname.endswith(".xls"):
        return extract_text_from_excel(file_bytes)
    else:
        # Fallback to UTF-8 text decoding
        try:
            return file_bytes.decode("utf-8", errors="ignore")
        except Exception:
            return ""


def summarize_context(text: str) -> str:
    """Use the LLM to summarize BOQ and other project specific details from the raw text."""
    if not text.strip():
        return ""
    
    llm = ChatOpenRouter(model="qwen/qwen3.7-flash", temperature=0)
    system_prompt = (
        "Kamu adalah asisten AI ahli yang bertugas mengekstrak informasi penting dari dokumen proyek "
        "(seperti Bill of Quantities/BOQ, spesifikasi material, dan ruang lingkup pekerjaan).\n\n"
        "Tugasmu: Buatlah ringkasan terstruktur dari teks yang diberikan. Fokus pada:\n"
        "1. Material utama yang digunakan\n"
        "2. Kuantitas atau volume pekerjaan\n"
        "3. Ruang lingkup spesifik proyek\n"
        "4. Persyaratan teknis khusus jika ada\n\n"
        "Ringkasan ini akan digunakan sebagai konteks tambahan untuk membuat dokumen RKS (Rencana Kerja dan Syarat-Syarat). "
        "Gunakan bahasa Indonesia yang baku dan jelas."
    )
    
    messages = [
        SystemMessage(content=system_prompt),
        HumanMessage(content=f"Teks dokumen:\n{text[:15000]}") # Limiting text length just in case
    ]
    
    response = llm.invoke(messages)
    return response.content


def extract_tender_metadata_from_docs(
    rks_bytes: Optional[bytes] = None,
    rks_filename: Optional[str] = None,
    boq_bytes: Optional[bytes] = None,
    boq_filename: Optional[str] = None,
    tanggal_str: str = "",
    resiko_csms: str = "Tinggi"
) -> Dict[str, Any]:
    """Extract procurement metadata and financial figures from RKS & BOQ files using LLM."""
    base_d = parse_date_safely(tanggal_str)
    formatted_tgl = format_id_date(base_d, with_day=False)
    schedules = calculate_tender_schedules(base_d)

    # Defaults
    defaults: Dict[str, Any] = {
        "nama_pengadaan": "PENGADAAN DAN PEMASANGAN FASILITAS OPERASIONAL REGIONAL KALIMANTAN",
        "nomor_tender": "No.Project/DT/PND970000/2026-S7",
        "tanggal": formatted_tgl,
        "prime_cost": 13283434534.0,
        "total_dengan_kr": 14346100000.0,
        "resiko_csms": resiko_csms or "Tinggi",
        "tkdn_minimal": 20.12,
        "metode_pemenuhan": "Tender Terbatas",
        "jenis_kontrak": "Gabungan Harga Satuan & Lumpsum",
        "pejabat_procurement_nama": "Rigga Widar Atmagi",
        "pejabat_procurement_jabatan": "Area Manager Procurement Kalimantan",
        "pejabat_berwenang": "Sr. Manager Opt. & Maint. Regional Kalimantan",
        "pengawas_pekerjaan": "Region Manager RPD Regional Kalimantan",
        "prebid_tanggal": schedules["prebid_tanggal"],
        "prebid_waktu": schedules["prebid_waktu"],
        "prebid_tempat": schedules["prebid_tempat"],
        "pemasukan_mulai": schedules["pemasukan_mulai"],
        "pemasukan_selesai": schedules["pemasukan_selesai"],
        "context_summary": "",
    }

    rks_text = ""
    if rks_bytes and rks_filename:
        print(f"Extracting RKS content from {rks_filename}...")
        rks_text = extract_text_from_file(rks_bytes, rks_filename)

    boq_text = ""
    if boq_bytes and boq_filename:
        print(f"Extracting BOQ content from {boq_filename}...")
        boq_text = extract_text_from_file(boq_bytes, boq_filename)

    # If no text was extracted from either document, return smart defaults immediately
    if not rks_text.strip() and not boq_text.strip():
        return defaults

    def _slice_doc_text(txt: str, max_chars: int = 14000) -> str:
        if not txt:
            return ""
        if len(txt) <= max_chars:
            return txt
        # Take beginning (title, scope, background) and end (approvals, pejabat, summary)
        half = max_chars // 2
        return txt[:half] + "\n\n[... potongan halaman tengah dilewati ...]\n\n" + txt[-half:]

    prompt = f"""
Kamu adalah Legal & Procurement Specialist PT Pertamina Patra Niaga.
Tugas kamu adalah menganalisis teks dokumen pendukung pengadaan (Dokumen RKS dan/atau Dokumen BOQ) dan mengekstrak parameter-parameter utama yang dibutuhkan untuk penyusunan Dokumen Tender resmi.

Input Parameter Awal:
- Tanggal Dokumen Tender: {formatted_tgl}
- Tingkat Risiko CSMS: {resiko_csms}

Panduan Ekstraksi:
1. "nama_pengadaan": Cari nama paket pekerjaan / judul pengadaan secara lengkap dan formal dari RKS atau BOQ. Contoh: "PENGADAAN DAN PEMASANGAN PIPA TRANSMISI MINYAK DAN GAS REGIONAL KALIMANTAN".
2. "prime_cost": Cari total biaya pokok riil pekerjaan (HPS Prime Cost / Subtotal BOQ sebelum Keuntungan & Risiko dan sebelum PPN) dalam bentuk angka float murni (contoh: 13283434534.0).
3. "total_dengan_kr": Cari total HPS yang sudah mencakup Keuntungan & Risiko (K&R). Jika di rekapitulasi BOQ hanya ada satu angka total biaya riil (Prime Cost) tanpa baris K&R terpisah, kalikan nilai prime_cost tersebut dengan 1.08 (K&R standar 8%).
4. "tkdn_minimal": Target persentase TKDN minimal jika tercantum dalam RKS/BOQ (contoh: 20.12). Jika tidak ada, gunakan default 20.12.
5. "nomor_tender": Nomor dokumen tender / RKS jika tercantum di dokumen. Jika tidak ditemukan, gunakan "No.Project/DT/PND970000/2026-S7".
6. "pejabat_procurement_nama" & "pejabat_procurement_jabatan": Nama & jabatan pejabat procurement jika tertera di RKS. Jika tidak ada, gunakan "Rigga Widar Atmagi" dan "Area Manager Procurement Kalimantan".
7. "pejabat_berwenang": Pejabat berwenang jika ada di RKS, default "Sr. Manager Opt. & Maint. Regional Kalimantan".
8. "pengawas_pekerjaan": Pengawas pekerjaan / direksi teknis jika ada di RKS, default "Region Manager RPD Regional Kalimantan".
9. "metode_pemenuhan": Metode pengadaan ("Tender Terbatas", "Tender Terbuka", "Pemilihan Langsung", atau "Penunjukan Langsung"). Default: "Tender Terbatas".
10. "jenis_kontrak": Jenis kontrak pengadaan ("Gabungan Harga Satuan & Lumpsum", "Harga Satuan", "Harga Lumpsum"). Default: "Gabungan Harga Satuan & Lumpsum".
11. "context_summary": Ringkasan terstruktur ruang lingkup pekerjaan (scope of work), spesifikasi material utama, dan kriteria kualifikasi teknis yang diperoleh dari RKS dan BOQ.

Isi Dokumen RKS:
{_slice_doc_text(rks_text, 14000) if rks_text else "(Dokumen RKS tidak dilampirkan)"}

Isi Dokumen BOQ:
{_slice_doc_text(boq_text, 14000) if boq_text else "(Dokumen BOQ tidak dilampirkan)"}

Format jawaban HARUS berupa JSON valid tanpa teks atau penjelasan lain di luar blok JSON:
{{
  "nama_pengadaan": "...",
  "prime_cost": 0.0,
  "total_dengan_kr": 0.0,
  "tkdn_minimal": 0.0,
  "nomor_tender": "...",
  "pejabat_procurement_nama": "...",
  "pejabat_procurement_jabatan": "...",
  "pejabat_berwenang": "...",
  "pengawas_pekerjaan": "...",
  "metode_pemenuhan": "...",
  "jenis_kontrak": "...",
  "context_summary": "..."
}}
"""

    try:
        response = llm.invoke([SystemMessage(content=prompt)])
        raw_text = response.content.strip()
        if raw_text.startswith("```"):
            raw_text = raw_text.split("```", 2)[1]
            if raw_text.startswith("json"):
                raw_text = raw_text[4:]
            raw_text = raw_text.rsplit("```", 1)[0]
        data = json.loads(raw_text.strip())

        # Merge with defaults & ensure types
        result = dict(defaults)
        for k, v in data.items():
            if v is not None and v != "":
                if k in ("prime_cost", "total_dengan_kr", "tkdn_minimal"):
                    try:
                        result[k] = float(v)
                    except Exception:
                        pass
                else:
                    result[k] = v

        # Safety calculation if total_dengan_kr or prime_cost are unbalanced
        if result["prime_cost"] > 0 and result["total_dengan_kr"] <= 0:
            result["total_dengan_kr"] = round(result["prime_cost"] * 1.08, 2)
        elif result["total_dengan_kr"] > 0 and result["prime_cost"] <= 0:
            result["prime_cost"] = round(result["total_dengan_kr"] / 1.08, 2)

        return result

    except Exception as e:
        print(f"Error extracting tender metadata with LLM: {e}")
        return defaults


